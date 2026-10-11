#!/usr/bin/env python3
"""
Assemble submission/ from the built manuscript. Run after scripts/build_manuscript_md.py.

  01_Main_Manuscript.docx        copy of paper/MANUSCRIPT.docx
  02_Cover_Letter.docx           generated here
  03_Figures/                    main Figures 1-6 (PNG + PDF), renamed Figure_1 ... Figure_6
  04_Supplementary_Material/     Supplementary Figures S1-S21 (PNG + PDF)
  05_Data_Code_Availability.txt  copied from the manuscript Declarations
  06_Declarations.txt            copied from the manuscript Declarations
  07_Blinded_Manuscript.docx     MANUSCRIPT.md with author identity removed, built by build_manuscript_docx.py
"""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
SUB = ROOT / "submission"
MD = PAPER / "MANUSCRIPT.md"
AUTHOR = "Zubair Md Ishaq"
EMAIL = "zubairishaq12@gmail.com"
GITHUB = "https://github.com/Zubair121Md/DR_RP"
IDENTITY = [AUTHOR, EMAIL, GITHUB, "Zubair121Md", "zubairishaq"]


def section(md: str, label: str) -> str:
    m = re.search(rf"^\*\*{re.escape(label)}\.\*\* (.+)$", md, flags=re.M)
    if not m:
        raise SystemExit(f"MISSING declaration in manuscript: {label}")
    return m.group(1).strip()


def blind(md: str) -> str:
    out = []
    for line in md.splitlines():
        if re.match(r"^\*\*(Author|Affiliation|Corresponding author):\*\*", line):
            continue
        if re.match(r"^\*\*(Author contributions \(CRediT\)|Acknowledgements)\.\*\*", line):
            continue
        if line.startswith("**Code availability.**"):
            line = ("**Code availability.** The training and evaluation code, configuration files, seed list, metric "
                    "files, environment record and table/figure scripts will be released in a public repository and "
                    "archived on Zenodo before publication; the link is withheld for anonymous review.")
        out.append(line)
    text = "\n".join(out)
    for token in IDENTITY:
        if token.lower() in text.lower():
            raise SystemExit(f"BLINDING_FAIL identity token remains: {token}")
    return text


def cover_letter(path: Path, title: str, code: str) -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(12)
    paras = [
        "11 October 2026",
        "Dear Editor,",
        f"Please consider the enclosed manuscript, “{title}”, for publication as an original research article.",
        "Problem addressed. Deep-learning systems for diabetic retinopathy (DR) grading often lose accuracy on "
        "images from datasets they were not trained on. Green-channel contrast-limited adaptive histogram "
        "equalisation (CLAHE) and threshold-based ordinal objectives are common design choices, but their "
        "behaviour under external testing is rarely reported, particularly for mild DR.",
        "Main contribution. We report an exploratory, seed-resolved evaluation of MobileViT-S trained only on "
        "APTOS 2019, crossing RGB or green-channel CLAHE input with cross-entropy or a mild-weighted shared-weight "
        "threshold objective, and evaluated on an internal APTOS test split and on IDRiD, Messidor-2 and DeepDRiD. "
        "Results are reported separately for two seed groups because a post-training audit found that seeds 42–44 "
        "received a truncated learning-rate schedule and seeds 45–47 the intended full cosine schedule. The audit "
        "also showed that the learned threshold head was not rank-consistent, so calibration is reported only for "
        "the cross-entropy models. An exact-duplicate audit of the APTOS partition and a schedule-stratified "
        "deduplication sensitivity analysis are included.",
        "Fit with the journal. The manuscript emphasises external validation, transparent disclosure of "
        "implementation deviations and reproducible reporting rather than a definitive performance claim, and "
        "should interest readers working on medical-image AI and DR screening.",
        "Exploratory and transparency-focused nature. The study is explicitly presented as exploratory. We do not "
        "claim that any configuration is generally preferable, we do not describe the threshold head as CORAL, and "
        "we do not report calibration for threshold-head models.",
        "Originality. The work is original, has not been published previously, and is not under consideration "
        "for publication elsewhere.",
        f"Data and code availability. All datasets are publicly available from their providers and are not "
        f"redistributed. {code}",
        "AI-assistance disclosure. AI-assisted tools were used to support code development and the drafting and "
        "editing of the manuscript. The author designed the study, reviewed all code, results and text, and takes "
        "full responsibility for the content.",
        "Competing interests. The author declares no conflicts of interest.",
        "Funding. This research received no external funding.",
        "Ethics. Ethics approval was not required; the study used only publicly available, de-identified datasets "
        "released for research.",
        "Thank you for considering this submission.",
        "Sincerely,",
        f"{AUTHOR}\nIndependent researcher\n{EMAIL}",
    ]
    for p in paras:
        doc.add_paragraph(p)
    doc.core_properties.author = AUTHOR
    doc.core_properties.title = "Cover letter"
    doc.save(str(path))


def main() -> None:
    md = MD.read_text()
    title = re.search(r"^# (.+)$", md, flags=re.M).group(1).strip()
    if SUB.exists():
        shutil.rmtree(SUB)
    (SUB / "03_Figures").mkdir(parents=True)
    (SUB / "04_Supplementary_Material").mkdir(parents=True)

    shutil.copy2(PAPER / "MANUSCRIPT.docx", SUB / "01_Main_Manuscript.docx")

    imgs = re.findall(r"^!\[(Figure (S?)(\d+))\..*?\]\((.+?)\)$", md, flags=re.M)
    for _, supp, n, rel in imgs:
        src = PAPER / rel
        dest_dir = SUB / ("04_Supplementary_Material" if supp else "03_Figures")
        stem = f"Figure_S{n}" if supp else f"Figure_{n}"
        for ext in (".png", ".pdf"):
            f = src.with_suffix(ext)
            if not f.is_file():
                raise SystemExit(f"MISSING figure file: {f.relative_to(ROOT)}")
            shutil.copy2(f, dest_dir / f"{stem}{ext}")
    print("COPIED figures:", sum(1 for x in imgs if not x[1]), "main,", sum(1 for x in imgs if x[1]), "supplementary")

    data, code = section(md, "Data availability"), section(md, "Code availability")
    (SUB / "05_Data_Code_Availability.txt").write_text(f"Data availability\n{data}\n\nCode availability\n{code}\n")
    decl = [f"{lab}. {section(md, lab)}" for lab in
            ("Funding", "Competing interests", "Author contributions (CRediT)", "Ethics approval",
             "Use of AI-assisted tools", "Acknowledgements")]
    (SUB / "06_Declarations.txt").write_text("\n\n".join(decl) + "\n")

    cover_letter(SUB / "02_Cover_Letter.docx", title, code)

    bmd = SUB / "07_Blinded_Manuscript.md"
    bmd.write_text(blind(md))
    subprocess.check_call([sys.executable, str(ROOT / "scripts" / "build_manuscript_docx.py"),
                           "--md", str(bmd), "--out", str(SUB / "07_Blinded_Manuscript.docx"),
                           "--author", "Anonymous"])
    bmd.unlink()
    doc = Document(str(SUB / "07_Blinded_Manuscript.docx"))
    full = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    leaks = [t for t in IDENTITY if t.lower() in full.lower()]
    if leaks or doc.core_properties.author != "Anonymous":
        raise SystemExit(f"BLINDING_FAIL in docx: {leaks} author={doc.core_properties.author!r}")
    print("WROTE submission/ (7 items)")
    print("BUILD_SUBMISSION_OK")


if __name__ == "__main__":
    main()
