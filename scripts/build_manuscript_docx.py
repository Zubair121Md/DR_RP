#!/usr/bin/env python3
"""
Convert paper/MANUSCRIPT.md (written by build_manuscript_md.py) into paper/MANUSCRIPT.docx.

Handles the subset of markdown used by the manuscript: headings, paragraphs with **bold**
and *italic*, numbered lists, pipe tables with a caption line above, images with a caption
paragraph below, display equations in \\[ ... \\] (rendered with matplotlib mathtext to
paper/figures/equations/), and the reference list.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.restricted import RESTRICTED_FIGURES  # noqa: E402
MD = ROOT / "paper" / "MANUSCRIPT.md"
DOCX = ROOT / "paper" / "MANUSCRIPT.docx"
EQ_DIR = ROOT / "paper" / "figures" / "equations"
FONT = "Times New Roman"
TEXT_WIDTH_IN = 6.27
IMG_MAX_IN = 6.2


def rel(p: Path) -> str:
    return str(p.relative_to(ROOT))


# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------

def set_font(style, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))

    normal = doc.styles["Normal"]
    set_font(normal, 11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    set_font(doc.styles["Title"], 16, bold=True)
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for bdr in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(bdr)
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    for name, size in (("Heading 1", 13), ("Heading 2", 11.5)):
        st = doc.styles[name]
        set_font(st, size, bold=True, italic=False)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    set_font(cap, 9.5, bold=False, italic=False)
    cap.paragraph_format.space_after = Pt(10)

    # Page numbers in the footer.
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    return doc


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def add_inline(par, text: str, size: float | None = None) -> None:
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            run.italic = True
        else:
            run = par.add_run(tok)
        if size:
            run.font.size = Pt(size)


def set_cell_borders(cell, **edges: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge, spec in edges.items():
        el = OxmlElement(f"w:{edge}")
        if spec == "none":
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), spec)
            el.set(qn("w:color"), "000000")
        borders.append(el)


def add_table(doc: Document, lines: list[str]) -> None:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    header, body = rows[0], [r for r in rows[2:]]
    ncol = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    size = 8.5 if ncol >= 5 else 9.5
    for ri, row in enumerate([header] + body):
        if len(row) != ncol:
            raise SystemExit(f"Table row has {len(row)} cells, header has {ncol}: {row}")
        for ci, text in enumerate(row):
            cell = table.cell(ri, ci)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.0
            if ri == 0:
                run = par.add_run(text)
                run.bold = True
                run.font.size = Pt(size)
            else:
                add_inline(par, text, size)
            edges = {"left": "none", "right": "none", "top": "none", "bottom": "none"}
            if ri == 0:
                edges["top"] = "8"
                edges["bottom"] = "4"
            if ri == len(body):
                edges["bottom"] = "8"
            set_cell_borders(cell, **edges)
    for r in table.rows:
        trpr = r._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def native_width_in(path: Path) -> float:
    with Image.open(path) as im:
        dpi = im.info.get("dpi", (600, 600))[0] or 600
        return im.width / float(dpi)


def add_image(doc: Document, path: Path, alt: str) -> None:
    if not path.is_file():
        print("MISSING", rel(path))
        if str(path.relative_to(ROOT / "paper")) in RESTRICTED_FIGURES:
            raise SystemExit(f"MISSING figure file {rel(path)}: it contains dataset fundus pixels and is not distributed "
                             "with the repository. Obtain the datasets (DATASET_PROVENANCE.md) and run "
                             "paper/figures/make_all.py to regenerate it before building the .docx.")
        raise SystemExit(f"MISSING figure file {rel(path)}")
    print("READ", rel(path))
    width = min(native_width_in(path), IMG_MAX_IN)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_after = Pt(2)
    run = par.add_run()
    run.add_picture(str(path), width=Inches(width))
    docpr = run._r.xpath(".//wp:docPr")
    if docpr:
        docpr[0].set("descr", alt)
        docpr[0].set("title", alt)


def render_equation(tex: str, n: int) -> Path:
    EQ_DIR.mkdir(parents=True, exist_ok=True)
    out = EQ_DIR / f"eq{n:02d}.png"
    fig = plt.figure(figsize=(0.01, 0.01))
    tex = tex.replace(r"\frac", r"\dfrac")
    fig.text(0, 0, f"${tex}$", fontsize=12, math_fontfamily="stix")
    fig.savefig(out, dpi=600, bbox_inches="tight", pad_inches=0.03, transparent=False, facecolor="white")
    plt.close(fig)
    print("WROTE", rel(out))
    return out


def add_equation(doc: Document, tex: str, n: int) -> None:
    path = render_equation(tex, n)
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
    par.add_run("\t")
    width = min(native_width_in(path), 5.0)
    run = par.add_run()
    run.add_picture(str(path), width=Inches(width))
    docpr = run._r.xpath(".//wp:docPr")
    if docpr:
        docpr[0].set("descr", f"Equation {n}")
    par.add_run(f"\t({n})")


# ---------------------------------------------------------------------------
# Main conversion
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", type=Path, default=MD)
    ap.add_argument("--out", type=Path, default=DOCX)
    ap.add_argument("--author", default="Zubair Md Ishaq", help="core-properties author; use 'Anonymous' for blinded builds")
    args = ap.parse_args()
    md, out = args.md.resolve(), args.out.resolve()
    if not md.is_file():
        print("MISSING", rel(md))
        raise SystemExit(1)
    print("READ", rel(md))
    text = re.sub(r"<!--.*?-->", "", md.read_text(), flags=re.S)
    lines = text.split("\n")
    doc = setup_document()

    i = 0
    eq_n = 0
    pending_caption_is_figure = False
    in_refs = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("# "):
            doc.add_paragraph(stripped[2:], style="Title")
            i += 1
            continue
        if stripped.startswith("## "):
            title = stripped[3:]
            if title.lower().startswith("supplementary"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            in_refs = title.lower() == "references"
            doc.add_heading(title, level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            i += 1
            continue

        if stripped == r"\[":
            j = i + 1
            buf = []
            while lines[j].strip() != r"\]":
                buf.append(lines[j].strip())
                j += 1
            eq_n += 1
            add_equation(doc, " ".join(buf), eq_n)
            i = j + 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if m:
            add_image(doc, ROOT / "paper" / m.group(2), m.group(1))
            pending_caption_is_figure = True
            i += 1
            continue

        if re.match(r"^\*\*(Figure S?\d+)\.\*\*", stripped):
            if not pending_caption_is_figure:
                raise SystemExit(f"Figure caption without an image above it: {stripped[:60]}")
            par = doc.add_paragraph(style="Caption")
            add_inline(par, stripped)
            pending_caption_is_figure = False
            i += 1
            continue

        if re.match(r"^\*\*Table S?\d+\.\*\*", stripped):
            par = doc.add_paragraph(style="Caption")
            par.paragraph_format.keep_with_next = True
            par.paragraph_format.space_after = Pt(3)
            add_inline(par, stripped)
            i += 1
            continue

        mnum = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if mnum:
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = Inches(0.3)
            pf.first_line_indent = Inches(-0.25)
            pf.space_after = Pt(3)
            add_inline(par, f"{mnum.group(1)}.\t{mnum.group(2)}")
            pf.tab_stops.add_tab_stop(Inches(0.3))
            i += 1
            continue

        par = doc.add_paragraph()
        if in_refs:
            pf = par.paragraph_format
            pf.left_indent = Inches(0.35)
            pf.first_line_indent = Inches(-0.35)
            pf.space_after = Pt(3)
            add_inline(par, stripped, size=9.5)
        elif stripped.startswith("*") and not stripped.startswith("**"):
            add_inline(par, stripped, size=9)
        else:
            add_inline(par, stripped)
        if line.endswith("  "):
            par.paragraph_format.space_after = Pt(0)
        i += 1

    doc.core_properties.title = lines[0].lstrip("# ").strip() if lines[0].startswith("#") else "Manuscript"
    doc.core_properties.author = args.author
    doc.core_properties.last_modified_by = args.author
    doc.core_properties.comments = ""
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print("WROTE", rel(out))
    print("BUILD_MANUSCRIPT_DOCX_OK")


if __name__ == "__main__":
    sys.exit(main())
